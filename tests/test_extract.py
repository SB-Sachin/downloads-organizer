import io
import wave
import zipfile

import pytest
from PIL import Image

from organizer.extract import extract

LIMIT = 2 * 1024**3


def _png_size(data: bytes) -> tuple[int, int]:
    assert data[:8] == b"\x89PNG\r\n\x1a\n"
    return Image.open(io.BytesIO(data)).size


def test_header_always_present(tmp_path):
    path = tmp_path / "Notes.TXT"
    path.write_text("hello")
    header = extract(path, LIMIT).header
    assert "Original filename: Notes.TXT" in header
    assert "Extension: .txt" in header
    assert "Size: 5 bytes" in header
    assert "Modified: 20" in header


@pytest.mark.parametrize("name", ["a.txt", "a.md", "a.csv", "a.json", "a.html", "a.log", "a.py"])
def test_text_types_first_4000_chars(tmp_path, name):
    path = tmp_path / name
    path.write_text("Cooking Merit Badge Requirements\n" + "x" * 10000)
    result = extract(path, LIMIT)
    assert result.kind == "text"
    assert len(result.text) == 4000
    assert result.text.startswith("Cooking Merit Badge")


def test_text_invalid_utf8_is_replaced(tmp_path):
    path = tmp_path / "bad.txt"
    path.write_bytes(b"caf\xe9 \xff\xfe ok")
    assert "�" in extract(path, LIMIT).text


def test_pdf_with_text(tmp_path):
    import pymupdf as fitz

    path = tmp_path / "bill.pdf"
    doc = fitz.open()
    for number in range(7):
        page = doc.new_page()
        page.insert_text((72, 72), f"Acme Power Electric Bill page {number + 1}\n" + "Amount due 132.18 " * 8)
    doc.save(path)
    result = extract(path, LIMIT)
    assert result.kind == "pdf text"
    assert "Acme Power Electric Bill page 1" in result.text
    assert "page 5" in result.text and "page 6" not in result.text  # first 5 pages only
    assert len(result.text) <= 6000
    assert not result.images


def test_scanned_pdf_renders_first_page(tmp_path):
    import pymupdf as fitz

    path = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page(width=612, height=792)  # no text layer
    doc.save(path)
    result = extract(path, LIMIT)
    assert result.kind.startswith("scanned pdf")
    assert result.text == ""
    width, height = _png_size(result.images[0])
    assert (width, height) == (850, 1100)  # 8.5 x 11 in at 100 dpi


@pytest.mark.parametrize("ext,fmt", [("png", "PNG"), ("jpg", "JPEG"), ("jpeg", "JPEG"),
                                     ("gif", "GIF"), ("webp", "WEBP"), ("heic", "HEIF")])
def test_images_resized_to_1024_png(tmp_path, ext, fmt):
    import pillow_heif

    pillow_heif.register_heif_opener()
    path = tmp_path / f"photo.{ext}"
    Image.new("RGB", (2048, 1000), (200, 30, 30)).save(path, format=fmt)
    result = extract(path, LIMIT)
    assert result.kind == "image"
    assert _png_size(result.images[0]) == (1024, 500)


def test_small_image_is_not_upscaled(tmp_path):
    path = tmp_path / "icon.png"
    Image.new("RGBA", (64, 32)).save(path)
    assert _png_size(extract(path, LIMIT).images[0]) == (64, 32)


def test_docx(tmp_path):
    import docx

    path = tmp_path / "essay.docx"
    document = docx.Document()
    document.add_paragraph("History Essay: The Erie Canal")
    for _ in range(100):
        document.add_paragraph("The canal opened in 1825. " * 5)
    document.save(path)
    result = extract(path, LIMIT)
    assert result.text.startswith("History Essay: The Erie Canal")
    assert len(result.text) == 4000


def test_xlsx(tmp_path):
    import openpyxl

    path = tmp_path / "budget.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    wb.create_sheet("Notes")
    ws.append(["Month", "Category", "Amount"])
    for row in range(1, 31):
        ws.append([f"month-{row}", "Food", row * 10])
    wb.save(path)
    text = extract(path, LIMIT).text
    assert "Sheets: Budget, Notes" in text
    assert "Header: Month | Category | Amount" in text
    assert "Row 10: month-10 | Food | 100" in text
    assert "month-11" not in text


def test_zip_lists_first_50_entries(tmp_path):
    path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(path, "w") as archive:
        for number in range(80):
            archive.writestr(f"src/file_{number:03}.py", "pass")
    text = extract(path, LIMIT).text
    assert "80 entries" in text
    assert "src/file_049.py" in text and "src/file_050.py" not in text


def test_audio_uses_mdls(tmp_path):
    path = tmp_path / "tone.wav"
    with wave.open(str(path), "wb") as wav:
        wav.setnchannels(1)
        wav.setsampwidth(2)
        wav.setframerate(8000)
        wav.writeframes(b"\x00\x00" * 8000)
    result = extract(path, LIMIT)
    # Spotlight may not have indexed a temp file, so only the route is asserted.
    assert result.kind == "audio/video metadata"
    assert "(null)" not in result.text
    assert not result.images


@pytest.mark.parametrize("name", ["Docker.dmg", "tool.pkg", "ubuntu.ISO", "Thing.app"])
def test_installers_skip_the_model(tmp_path, name):
    path = tmp_path / name
    path.write_bytes(b"\x00" * 64)
    result = extract(path, LIMIT)
    assert result.skip_model and result.kind == "installer"


def test_unknown_type_uses_file_command(tmp_path):
    path = tmp_path / "mystery.xyz"
    path.write_bytes(b"\x7fELF\x02\x01\x01" + b"\x00" * 100)
    result = extract(path, LIMIT)
    assert result.kind == "file type description"
    assert "ELF" in result.text


def test_over_size_limit_uses_metadata_only(tmp_path):
    path = tmp_path / "big.txt"
    path.write_text("secret content " * 100)
    result = extract(path, max_content_bytes=100)
    assert result.text == "" and not result.images
    assert "over size limit" in result.kind
    assert "big.txt" in result.prompt()


def test_corrupt_file_falls_back_to_description(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"this is not a zip container")
    result = extract(path, LIMIT)
    assert result.kind == "unreadable content"
    assert "could not be read" in result.text
