from __future__ import annotations

import io
import os
import subprocess
import zipfile
from dataclasses import dataclass, field
from datetime import datetime

TEXT_CHARS = 4000
PDF_CHARS = 6000
PDF_PAGES = 5
PDF_SCANNED_THRESHOLD = 100
IMAGE_MAX_SIDE = 1024

TEXT_EXTS = {
    ".txt", ".md", ".markdown", ".rst", ".csv", ".tsv", ".json", ".jsonl", ".html", ".htm",
    ".xml", ".log", ".yaml", ".yml", ".toml", ".ini", ".cfg", ".conf", ".env", ".sql",
    ".py", ".js", ".mjs", ".ts", ".jsx", ".tsx", ".css", ".scss", ".c", ".h", ".cpp", ".hpp",
    ".cc", ".m", ".java", ".kt", ".go", ".rs", ".rb", ".php", ".swift", ".sh", ".zsh",
    ".bash", ".lua", ".pl", ".r", ".tex", ".srt", ".ics", ".vcf", ".gcode", ".scad",
}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".heic"}
AUDIO_VIDEO_EXTS = {
    ".mp3", ".m4a", ".aac", ".wav", ".flac", ".aiff", ".aif", ".ogg", ".opus", ".wma",
    ".mp4", ".m4v", ".mov", ".mkv", ".avi", ".webm", ".wmv", ".mpg", ".mpeg",
}
INSTALLER_EXTS = {".dmg", ".pkg", ".app", ".iso"}
MDLS_FIELDS = ("kMDItemTitle", "kMDItemDurationSeconds", "kMDItemAuthors")


@dataclass
class Extraction:
    kind: str
    header: str
    text: str = ""
    images: list[bytes] = field(default_factory=list)
    skip_model: bool = False  # installers: name is cleaned locally, no model call

    def prompt(self) -> str:
        if self.text:
            return f"{self.header}\n\nContent ({self.kind}):\n{self.text}"
        if self.images:
            return f"{self.header}\n\nContent ({self.kind}): see the attached image."
        return f"{self.header}\n\nNo content available ({self.kind}). Use the filename and metadata."


def _human_size(size: int) -> str:
    value = float(size)
    for unit in ("bytes", "KB", "MB", "GB"):
        if value < 1024 or unit == "GB":
            return f"{value:.0f} {unit}" if unit == "bytes" else f"{value:.1f} {unit}"
        value /= 1024
    return f"{size} bytes"


def _header(path: str, st: os.stat_result) -> str:
    name = os.path.basename(path)
    ext = os.path.splitext(name)[1].lower() or "(none)"
    modified = datetime.fromtimestamp(st.st_mtime).strftime("%Y-%m-%d %H:%M")
    return (
        f"Original filename: {name}\n"
        f"Extension: {ext}\n"
        f"Size: {_human_size(st.st_size)}\n"
        f"Modified: {modified}"
    )


def _text(path: str) -> str:
    # 4 bytes per character is the UTF-8 worst case, so this always yields enough.
    with open(path, "rb") as fh:
        raw = fh.read(TEXT_CHARS * 4)
    return raw.decode("utf-8", errors="replace")[:TEXT_CHARS]


def _pdf(path: str) -> tuple[str, str, list[bytes]]:
    import pymupdf as fitz

    with fitz.open(path) as doc:
        if doc.needs_pass:
            return "password-protected pdf", "", []
        parts: list[str] = []
        total = 0
        for index in range(min(PDF_PAGES, doc.page_count)):
            chunk = doc[index].get_text()
            parts.append(chunk)
            total += len(chunk)
            if total >= PDF_CHARS:
                break
        text = "\n".join(parts).strip()[:PDF_CHARS]
        if len(text) >= PDF_SCANNED_THRESHOLD or doc.page_count == 0:
            return "pdf text", text, []
        png = doc[0].get_pixmap(dpi=100).tobytes("png")
        return "scanned pdf, page 1", "", [png]


def _image(path: str) -> bytes:
    from PIL import Image
    import pillow_heif

    pillow_heif.register_heif_opener()
    with Image.open(path) as img:
        img.seek(0)  # first frame of animated GIF/WebP
        frame = img.convert("RGBA" if "A" in img.getbands() or "transparency" in img.info else "RGB")
        frame.thumbnail((IMAGE_MAX_SIDE, IMAGE_MAX_SIDE))
        buffer = io.BytesIO()
        frame.save(buffer, format="PNG")
        return buffer.getvalue()


def _docx(path: str) -> str:
    import docx

    out: list[str] = []
    total = 0
    for para in docx.Document(path).paragraphs:
        if para.text.strip():
            out.append(para.text)
            total += len(para.text) + 1
            if total >= TEXT_CHARS:
                break
    return "\n".join(out)[:TEXT_CHARS]


def _xlsx(path: str) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        lines = ["Sheets: " + ", ".join(wb.sheetnames)]
        sheet = wb[wb.sheetnames[0]]
        # Row 1 is the header, then the first 10 data rows.
        for number, row in enumerate(sheet.iter_rows(min_row=1, max_row=11, max_col=30, values_only=True)):
            cells = ["" if cell is None else str(cell) for cell in row]
            label = "Header" if number == 0 else f"Row {number}"
            lines.append(f"{label}: " + " | ".join(cells).rstrip(" |"))
        return "\n".join(lines)[:TEXT_CHARS]
    finally:
        wb.close()


def _zip(path: str) -> str:
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
    listing = "\n".join(names[:50])
    return f"{len(names)} entries. First {min(50, len(names))}:\n{listing}"


def _run(args: list[str], timeout: float = 15) -> str:
    try:
        result = subprocess.run(args, capture_output=True, text=True, timeout=timeout, errors="replace")
    except (OSError, subprocess.SubprocessError):
        return ""
    return result.stdout.strip()


def _mdls(path: str) -> str:
    args = ["mdls"]
    for name in MDLS_FIELDS:
        args += ["-name", name]
    output = _run(args + [path])
    kept = [line for line in output.splitlines() if "(null)" not in line]
    return "\n".join(kept)


def _file_command(path: str) -> str:
    return _run(["file", "-b", path])[:1000]


def extract(path: str | os.PathLike, max_content_bytes: int) -> Extraction:
    """Build the model input for one file. Never raises for unreadable *content*:
    a corrupt file falls back to metadata plus a note."""
    path = os.fspath(path)
    st = os.stat(path)
    header = _header(path, st)
    ext = os.path.splitext(path)[1].lower()

    if ext in INSTALLER_EXTS:
        return Extraction("installer", header, skip_model=True)
    if st.st_size > max_content_bytes:
        return Extraction("metadata only, file over size limit", header)

    try:
        if ext in TEXT_EXTS:
            return Extraction("text", header, text=_text(path))
        if ext == ".pdf":
            kind, text, images = _pdf(path)
            return Extraction(kind, header, text=text, images=images)
        if ext in IMAGE_EXTS:
            return Extraction("image", header, images=[_image(path)])
        if ext == ".docx":
            return Extraction("word document text", header, text=_docx(path))
        if ext == ".xlsx":
            return Extraction("spreadsheet", header, text=_xlsx(path))
        if ext == ".zip":
            return Extraction("zip archive listing", header, text=_zip(path))
        if ext in AUDIO_VIDEO_EXTS:
            return Extraction("audio/video metadata", header, text=_mdls(path))
    except FileNotFoundError:
        raise
    except Exception as exc:  # corrupt or mislabelled file: describe it instead
        note = f"(content could not be read: {type(exc).__name__})"
        return Extraction("unreadable content", header, text=f"{note}\n{_file_command(path)}".strip())
    return Extraction("file type description", header, text=_file_command(path))
